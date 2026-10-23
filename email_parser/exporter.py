"""Read/write JSON, CSV and (formatted) XLSX artefacts."""

from __future__ import annotations

import csv
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from .models import (
    SOURCE_FB_PAGE,
    SOURCE_FB_SNIPPET,
    SOURCE_GUESS,
    SOURCE_LABELS,
    SOURCE_OSM,
    SOURCE_SEARCH,
    SOURCE_SITE,
    SOURCE_SITE_JS,
    SOURCE_STATUS,
    STATUS_NOT_FOUND,
    ContactResult,
    Restaurant,
)
from .utils import Area, ensure_parent, unique

CSV_COLUMNS = (
    ("Name", "title"),
    ("Category", "category"),
    ("Rating", "rating"),
    ("Reviews", "reviews"),
    ("Price", "price"),
    ("Website", "website"),
    ("Emails", "emails"),
    ("Primary Email", "primary_email"),
    ("Facebook", "facebook_url"),
    ("Phone", "phone"),
    ("Address", "address"),
    ("Status", "status"),
    ("Sources", "sources"),
    ("Google Maps URL", "google_maps_url"),
    ("Note", "note"),
    ("Latitude", "latitude"),
    ("Longitude", "longitude"),
    ("Query", "source_query"),
    ("Primary Email Source", "primary_email_source_label_en"),
    ("Email Sources", "email_sources_text"),
)

SOURCE_CSV_LABELS = {
    SOURCE_SITE: "website",
    SOURCE_SITE_JS: "website (js render)",
    SOURCE_FB_PAGE: "facebook (page)",
    SOURCE_FB_SNIPPET: "facebook (search snippet)",
    SOURCE_SEARCH: "web search",
    SOURCE_OSM: "OpenStreetMap",
    SOURCE_GUESS: "guessed MX",
}

STATUS_LABELS = {
    "Found on site": "Найден на сайте",
    "Found on FB": "Найден на Facebook",
    "Found via web search": "Найден в поиске",
    "Found via JS render": "Найден после рендера",
    "Found on OpenStreetMap": "Найден в OpenStreetMap",
    "Guessed (unverified)": "Предположение (не проверено)",
    "Not found": "Не найден",
}

STATUS_FILLS = {
    "Found on site": "C6EFCE",
    "Found on FB": "BDD7EE",
    "Found via web search": "DDEBF7",
    "Found via JS render": "E2EFDA",
    "Found on OpenStreetMap": "D9E1F2",
    "Guessed (unverified)": "FFEB9C",
    "Not found": "FFC7CE",
}


def save_restaurants_json(path: str | Path, items: list[Restaurant]) -> Path:
    path = ensure_parent(path)
    payload = [item.to_dict() for item in items]
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def load_restaurants(path: str | Path) -> list[Restaurant]:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(
            f"{path} not found. Run the maps stage first: "
            f'python main.py -q "restaurants in City, ST" --stage maps'
        )
    data = json.loads(path.read_text(encoding="utf-8"))
    return [Restaurant.from_dict(row) for row in data]


def save_results_json(path: str | Path, results: list[ContactResult]) -> Path:
    path = ensure_parent(path)
    payload = [row.to_dict() for row in results]
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def save_results_csv(path: str | Path, results: list[ContactResult]) -> Path:
    path = ensure_parent(path)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=[name for name, _ in CSV_COLUMNS])
        writer.writeheader()
        for row in results:
            data = row.to_dict()
            data["primary_email"] = row.primary_email
            data["primary_email_source_label_en"] = SOURCE_CSV_LABELS.get(
                row.primary_email_source, row.primary_email_source
            )
            data["email_sources_text"] = "; ".join(
                f"{email} <- {SOURCE_CSV_LABELS.get(source, source)}"
                for email, source in row.email_sources.items()
            )
            writer.writerow({name: data.get(key, "") for name, key in CSV_COLUMNS})
    return path


# --------------------------------------------------------------------- XLSX
XLSX_HEADERS = (
    ("№", 5),
    ("Название", 32),
    ("Категория", 22),
    ("Рейтинг", 8),
    ("Отзывов", 9),
    ("Цена", 7),
    ("Сайт", 34),
    ("E-mail", 32),
    ("Основной e-mail", 26),
    ("Facebook", 34),
    ("Телефон", 16),
    ("Адрес", 40),
    ("Статус", 24),
    ("Источник", 22),
    ("Google Maps", 24),
    ("Примечание", 30),
    ("Широта", 11),
    ("Долгота", 11),
    ("Запрос", 30),
    ("Откуда основной e-mail", 26),
    ("Источники e-mail", 48),
)


def save_results_xlsx(
    path: str | Path,
    results: list[ContactResult],
    meta: dict[str, Any] | None = None,
) -> Path:
    """Write a human-friendly report with clickable links and a summary sheet."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - depends on the environment
        raise RuntimeError(
            "openpyxl is required for the Excel report: pip install openpyxl"
        ) from exc

    meta = meta or {}
    path = ensure_parent(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Контакты"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    link_font = Font(color="0563C1", underline="single")
    wrap = Alignment(vertical="top", wrap_text=True)
    center = Alignment(horizontal="center", vertical="top")

    for index, (name, width) in enumerate(XLSX_HEADERS, start=1):
        cell = sheet.cell(row=1, column=index, value=name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 28

    for row_index, row in enumerate(results, start=2):
        status_label = STATUS_LABELS.get(row.status, row.status)
        values = [
            row_index - 1,
            row.title,
            row.category,
            row.rating,
            row.reviews,
            row.price,
            row.website,
            ", ".join(row.emails),
            row.primary_email,
            row.facebook_url,
            row.phone,
            row.address,
            status_label,
            "; ".join(row.sources),
            row.google_maps_url,
            row.note,
            row.latitude,
            row.longitude,
            row.source_query,
            row.primary_email_source_label,
            "; ".join(
                f"{email} ← {SOURCE_LABELS.get(source, source)}"
                for email, source in row.email_sources.items()
            ),
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column, value=value)
            cell.alignment = wrap
        for column in (1, 4, 5, 6, 13, 17, 18):
            sheet.cell(row=row_index, column=column).alignment = center

        for column, url in ((7, row.website), (10, row.facebook_url), (15, row.google_maps_url)):
            cell = sheet.cell(row=row_index, column=column)
            if url:
                cell.hyperlink = url
                cell.font = link_font

        primary = sheet.cell(row=row_index, column=9)
        if row.primary_email:
            primary.hyperlink = f"mailto:{row.primary_email}"
            primary.font = link_font

        status_cell = sheet.cell(row=row_index, column=13)
        fill = STATUS_FILLS.get(row.status)
        if fill:
            status_cell.fill = PatternFill("solid", fgColor=fill)

    last_row = max(2, len(results) + 1)
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = f"A1:U{last_row}"

    _write_summary(workbook, results, meta)
    workbook.save(path)
    workbook.close()
    return path


def _write_summary(workbook, results: list[ContactResult], meta: dict[str, Any]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    sheet = workbook.create_sheet("Сводка")
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 90

    title = sheet.cell(row=1, column=1, value="Отчёт email-parser")
    title.font = Font(bold=True, size=15)
    title.fill = PatternFill("solid", fgColor="DDEBF7")

    generated = meta.get("generated") or datetime.now().strftime("%Y-%m-%d %H:%M")
    queries = meta.get("queries") or ([meta["query"]] if meta.get("query") else [])
    if not queries:
        queries = unique(row.source_query for row in results if row.source_query)
    categories = meta.get("categories") or []
    if not categories:
        categories = unique(row.category for row in results if row.category)
    area = meta.get("area")
    area_text, map_link = "", ""
    if isinstance(area, Area):
        area_text = area.as_text()
        lat, lng = area.center
        map_link = f"https://www.google.com/maps/@{lat:.6f},{lng:.6f},14z"
    elif area:
        area_text = str(area)

    stats = meta.get("stats") or {}
    rows: list[tuple[str, Any]] = [
        ("Сформирован", generated),
        ("Запросы", "\n".join(queries) if queries else "—"),
        ("Область (south,west,north,east)", area_text or "вся выдача по запросу"),
        ("Провайдер Maps", meta.get("provider") or "—"),
        ("Категорий", len(meta.get("categories") or []) or "—"),
        ("Заведений в отчёте", len(results)),
        ("С e-mail", sum(1 for row in results if row.emails)),
        ("Без e-mail", sum(1 for row in results if not row.emails)),
    ]
    if stats:
        rows.extend(
            [
                ("Найдено на сайте", stats.get("from_site", 0)),
                ("Найдено на Facebook", stats.get("from_facebook", 0)),
                ("Найдено в поиске", stats.get("from_search", 0)),
                ("Найдено рендером", stats.get("from_js", 0)),
                ("Предположений", stats.get("from_guess", 0)),
                ("Ошибок", stats.get("errors", 0)),
                ("Длительность, сек", round(float(stats.get("duration") or 0), 1)),
            ]
        )

    current = 3
    for label, value in rows:
        name_cell = sheet.cell(row=current, column=1, value=label)
        name_cell.font = Font(bold=True)
        value_cell = sheet.cell(row=current, column=2, value=value)
        value_cell.alignment = Alignment(vertical="top", wrap_text=True)
        if label.startswith("Область") and map_link:
            value_cell.hyperlink = map_link
            value_cell.font = Font(color="0563C1", underline="single")
        current += 1

    current += 1
    legend_title = sheet.cell(row=current, column=1, value="Статусы")
    legend_title.font = Font(bold=True, size=12)
    current += 1
    for status, label in STATUS_LABELS.items():
        cell = sheet.cell(row=current, column=1, value=label)
        fill = STATUS_FILLS.get(status)
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        sheet.cell(row=current, column=2, value=status)
        current += 1

    current += 1
    note = sheet.cell(row=current, column=1, value="Как проверить")
    note.font = Font(bold=True, size=12)
    current += 1
    explanation = (
        "Для каждой строки откройте «Сайт», «Facebook» или «Google Maps» (ссылки "
        "кликабельны) и убедитесь, что адрес действительно принадлежит заведению. "
        "Статус «Предположение» — это НЕ найденный адрес, а догадка по MX-записям: "
        "её нужно проверять вручную. Пустой сайт означает, что у заведения нет "
        "сайта в Google Maps — почта искалась в Facebook и поиске."
    )
    sheet.cell(row=current, column=1, value=explanation).alignment = Alignment(
        vertical="top", wrap_text=True
    )
    sheet.merge_cells(start_row=current, start_column=1, end_row=current, end_column=2)
    sheet.row_dimensions[current].height = 90


STATUS_VALUES = {label: value for value, label in STATUS_LABELS.items()}
SOURCE_VALUES = {label: value for value, label in SOURCE_LABELS.items()}
SOURCE_EN_VALUES = {label: value for value, label in SOURCE_CSV_LABELS.items()}


def _source_from_status(status: str) -> str:
    lowered = (status or "").lower()
    if "openstreetmap" in lowered or "open street map" in lowered or "osm" in lowered:
        return SOURCE_OSM
    if "facebook" in lowered or "fb" in lowered:
        return SOURCE_FB_PAGE
    if "рендер" in lowered or "render" in lowered:
        return SOURCE_SITE_JS
    if "поиск" in lowered or "search" in lowered:
        return SOURCE_SEARCH
    if "предполож" in lowered or "guess" in lowered:
        return SOURCE_GUESS
    if "сайт" in lowered or "site" in lowered:
        return SOURCE_SITE
    return ""


def _as_float(value) -> float | None:
    try:
        return float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return None


def _as_int(value) -> int | None:
    try:
        return int(float(str(value).replace(",", ".")))
    except (TypeError, ValueError):
        return None


def _split_emails(value) -> list[str]:
    return [
        item.strip().lower()
        for item in re.split(r"[,;]\s*", str(value or ""))
        if "@" in item and item.strip()
    ]


def load_results_json(path: str | Path) -> list[ContactResult]:
    path = Path(path)
    payload = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(payload, dict):
        payload = payload.get("results") or payload.get("contacts") or []
    return [ContactResult.from_dict(row) for row in payload]


def load_results_xlsx(path: str | Path) -> list[ContactResult]:
    """Read the «Контакты» sheet back, e.g. to send the application e-mails."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depends on the environment
        raise RuntimeError(
            "openpyxl is required to read the Excel report: pip install openpyxl"
        ) from exc

    path = Path(path)
    workbook = load_workbook(filename=path, read_only=True, data_only=True)
    sheet = workbook["Контакты"] if "Контакты" in workbook.sheetnames else workbook.active
    # Read once and close immediately: on Windows an open read-only workbook
    # keeps the .xlsx locked until the process exits.
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    if not rows:
        return []
    headers = rows[0]

    index = {
        str(value).strip(): position
        for position, value in enumerate(headers)
        if value is not None
    }

    def cell(values, *names):
        for name in names:
            position = index.get(name)
            if position is not None and position < len(values):
                return values[position]
        return None

    results: list[ContactResult] = []
    for values in rows[1:]:
        title = str(cell(values, "Название", "Name") or "").strip()
        if not title:
            continue

        emails = _split_emails(cell(values, "E-mail", "Emails"))
        primary = str(
            cell(values, "Основной e-mail", "Primary Email") or ""
        ).strip().lower()
        if primary and "@" in primary:
            emails = [primary] + [email for email in emails if email != primary]

        source_label = str(
            cell(
                values,
                "Откуда основной e-mail",
                "Primary Email Source",
            )
            or ""
        ).strip()
        source = (
            SOURCE_VALUES.get(source_label)
            or SOURCE_EN_VALUES.get(source_label)
            or ""
        )
        status_label = str(cell(values, "Статус", "Status") or "").strip()
        status = STATUS_VALUES.get(status_label, status_label)
        if status not in STATUS_LABELS:
            status = ""
        if not source:
            source = _source_from_status(status_label)
        if not status:
            status = SOURCE_STATUS.get(source, STATUS_NOT_FOUND)

        sources_raw = str(cell(values, "Источник", "Sources") or "")
        sources = [item.strip() for item in sources_raw.split(";") if item.strip()]

        results.append(
            ContactResult(
                title=title,
                website=str(cell(values, "Сайт", "Website") or "").strip(),
                emails=emails,
                facebook_url=str(
                    cell(values, "Facebook") or ""
                ).strip(),
                phone=str(cell(values, "Телефон", "Phone") or "").strip(),
                address=str(cell(values, "Адрес", "Address") or "").strip(),
                status=status,
                sources=sources,
                google_maps_url=str(
                    cell(values, "Google Maps", "Google Maps URL") or ""
                ).strip(),
                note=str(cell(values, "Примечание", "Note") or "").strip(),
                category=str(cell(values, "Категория", "Category") or "").strip(),
                rating=_as_float(cell(values, "Рейтинг", "Rating")),
                reviews=_as_int(cell(values, "Отзывов", "Reviews")),
                price=str(cell(values, "Цена", "Price") or "").strip(),
                latitude=_as_float(cell(values, "Широта", "Latitude")),
                longitude=_as_float(cell(values, "Долгота", "Longitude")),
                source_query=str(cell(values, "Запрос", "Query") or "").strip(),
                email_sources={primary: source} if primary and source else {},
                primary_email_override=primary,
            )
        )
    return results


def load_contacts(path: str | Path) -> list[ContactResult]:
    """Load an existing report (xlsx preferred, JSON also accepted)."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"{path} not found")
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return load_results_xlsx(path)
    if suffix == ".json":
        return load_results_json(path)
    raise ValueError(f"unsupported report format: {path.suffix} (use .xlsx or .json)")


def rows_with_email(results: Iterable[ContactResult]) -> list[ContactResult]:
    return [row for row in results if row.emails]
