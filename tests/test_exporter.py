"""Excel report contents."""

from __future__ import annotations

from openpyxl import load_workbook

from email_parser.exporter import load_results_xlsx, save_results_xlsx
from email_parser.models import (
    STATUS_CHAIN,
    STATUS_FB,
    STATUS_GUESS,
    STATUS_NOT_FOUND,
    STATUS_SITE,
    ContactResult,
)
from email_parser.utils import Area


def make_results() -> list[ContactResult]:
    return [
        ContactResult(
            title="Coastal Salt",
            website="https://coastalsaltoc.com",
            emails=["info@coastalsaltoc.com"],
            status=STATUS_SITE,
            sources=["website"],
            phone="(410) 973-7258",
            address="1601 Atlantic Ave",
            google_maps_url="https://www.google.com/maps/place/?q=place_id:abc",
            category="New American restaurant",
            rating=4.5,
            reviews=1200,
            latitude=38.33,
            longitude=-75.08,
            source_query="restaurant in Ocean City, MD",
        ),
        ContactResult(
            title="Shawarma Stand",
            facebook_url="https://www.facebook.com/shawarmastand",
            emails=["order@shawarma.example"],
            status=STATUS_FB,
            sources=["facebook (page render)"],
        ),
        ContactResult(title="Kiosk", status=STATUS_NOT_FOUND, note="нет сайта"),
        ContactResult(
            title="Guessed Place",
            website="https://guessed.example",
            emails=["info@guessed.example"],
            status=STATUS_GUESS,
            note="unverified",
        ),
    ]


def test_xlsx_has_links_filters_and_summary(tmp_path):
    path = tmp_path / "report.xlsx"
    area = Area(south=38.30, west=-75.10, north=38.40, east=-75.03)
    save_results_xlsx(
        path,
        make_results(),
        {
            "query": "Ocean City, MD",
            "queries": ["restaurant in Ocean City, MD", "shawarma in Ocean City, MD"],
            "area": area,
            "provider": "apify",
            "categories": ["restaurant", "shawarma"],
            "stats": {"from_site": 1, "from_facebook": 1, "from_guess": 1},
        },
    )
    assert path.exists()

    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Контакты", "Сводка"]

    sheet = workbook["Контакты"]
    assert sheet["A1"].value == "№"
    assert sheet["B2"].value == "Coastal Salt"
    assert sheet["G2"].hyperlink.target == "https://coastalsaltoc.com"
    assert sheet["I2"].hyperlink.target == "mailto:info@coastalsaltoc.com"
    assert sheet["J3"].value == "https://www.facebook.com/shawarmastand"
    assert sheet["M2"].value == "Найден на сайте"
    assert sheet.freeze_panes == "B2"
    assert sheet.auto_filter.ref == "A1:V5"
    assert sheet.max_row == 5
    assert sheet["T1"].value == "Откуда основной e-mail"
    assert sheet["V1"].value == "Сеть"

    summary = workbook["Сводка"]
    values = [cell.value for row in summary.iter_rows() for cell in row]
    assert "38.300000,-75.100000,38.400000,-75.030000" in values
    assert any(
        isinstance(value, str) and "Предположение" in value for value in values
    )


def test_chain_rows_get_separate_sheet_and_roundtrip(tmp_path):
    path = tmp_path / "chains.xlsx"
    save_results_xlsx(
        path,
        [
            ContactResult(
                title="Dunkin'",
                website="https://www.dunkindonuts.com/",
                chain="Dunkin'",
                status=STATUS_CHAIN,
            ),
            ContactResult(
                title="Coastal Salt",
                emails=["info@coastalsaltoc.com"],
                status=STATUS_SITE,
            ),
        ],
    )
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Контакты", "Сети", "Сводка"]
    assert workbook["Контакты"]["V2"].value == "Dunkin'"
    assert workbook["Сети"]["A2"].value == "Dunkin'"
    assert workbook["Сети"]["G2"].value == "Сеть — заявка на сайте"

    loaded = load_results_xlsx(path)
    assert loaded[0].chain == "Dunkin'"
    assert loaded[0].status == STATUS_CHAIN
    assert loaded[0].emails == []


def test_csv_contains_new_columns(tmp_path):
    from email_parser.exporter import save_results_csv

    path = tmp_path / "report.csv"
    save_results_csv(path, make_results())
    header = path.read_text(encoding="utf-8-sig").splitlines()[0]
    for column in ("Category", "Rating", "Facebook", "Latitude", "Query"):
        assert column in header
